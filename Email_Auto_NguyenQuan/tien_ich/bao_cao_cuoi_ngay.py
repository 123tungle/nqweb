"""Gửi email báo cáo bán hàng cuối ngày (PDF / Excel) cho Sếp.

Logic canh giờ / Chủ nhật / chống gửi trùng nằm trong store SQL (NQW_BaoCaoEmail_CanGuiStp),
nên SQL Agent chỉ cần chạy file này mỗi ~10 phút; tới giờ đã đặt trong bảng NQW_BaoCaoEmailCf
thì script mới thực sự gửi (1 lần/ngày).

Chạy tay:
  python tien_ich/bao_cao_cuoi_ngay.py                 # canh giờ theo cấu hình
  python tien_ich/bao_cao_cuoi_ngay.py --force         # gửi ngay, bỏ qua canh giờ (đánh dấu đã gửi hôm nay)
  python tien_ich/bao_cao_cuoi_ngay.py --ngay 2025-04-05 --preview   # chỉ tạo file, KHÔNG gửi
  python tien_ich/bao_cao_cuoi_ngay.py --test-email    # gửi 1 mail thử để kiểm tra SMTP
"""
from __future__ import annotations

import argparse
import datetime as dt
import smtplib
import ssl
import sys
from email.message import EmailMessage
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE))

from app import config          # noqa: E402
from app.erp import Erp         # noqa: E402

FONT_DIR = Path(__file__).resolve().parent / "fonts"
OUT_DIR = BASE / "data" / "bao_cao"


# ----------------------------------------------------------------------------- tiện ích
def tien(n) -> str:
    """Định dạng số tiền kiểu Việt Nam: 1.679.807.433"""
    try:
        return f"{int(round(float(n or 0))):,}".replace(",", ".")
    except (TypeError, ValueError):
        return "0"


def tach_email(s: str) -> list[str]:
    if not s:
        return []
    for sep in (";", "\n", " "):
        s = s.replace(sep, ",")
    return [e.strip() for e in s.split(",") if e.strip()]


def dinh_dang_ngay(d) -> str:
    if isinstance(d, (dt.date, dt.datetime)):
        return d.strftime("%d/%m/%Y")
    return str(d)


# ----------------------------------------------------------------------------- gọi store
def _dict1(cur):
    while cur.description is None:
        if not cur.nextset():
            return None
    cols = [c[0] for c in cur.description]
    row = cur.fetchone()
    return dict(zip(cols, row)) if row else {}


def kiem_tra(erp: Erp, now: dt.datetime | None = None) -> dict:
    raw = erp.engine.raw_connection()
    try:
        cur = raw.cursor()
        if now is not None:
            cur.execute("SET NOCOUNT ON; EXEC dbo.NQW_BaoCaoEmail_CanGuiStp @Now = ?", [now])
        else:
            cur.execute("SET NOCOUNT ON; EXEC dbo.NQW_BaoCaoEmail_CanGuiStp")
        return _dict1(cur) or {}
    finally:
        raw.rollback()
        raw.close()


def lay_du_lieu(erp: Erp, ngay) -> tuple[dict, list[dict]]:
    raw = erp.engine.raw_connection()
    try:
        cur = raw.cursor()
        cur.execute("SET NOCOUNT ON; EXEC dbo.NQW_BaoCaoCuoiNgay_DuLieuStp @Ngay = ?", [str(ngay)])
        while cur.description is None:
            if not cur.nextset():
                return {}, []
        cols1 = [c[0] for c in cur.description]
        r1 = cur.fetchone()
        tongquan = dict(zip(cols1, r1)) if r1 else {}
        theonv: list[dict] = []
        if cur.nextset() and cur.description is not None:
            cols2 = [c[0] for c in cur.description]
            theonv = [dict(zip(cols2, r)) for r in cur.fetchall()]
        return tongquan, theonv
    finally:
        raw.rollback()
        raw.close()


def danh_dau_da_gui(erp: Erp, ngay_gui: dt.date) -> None:
    raw = erp.engine.raw_connection()
    try:
        cur = raw.cursor()
        cur.execute("SET NOCOUNT ON; EXEC dbo.NQW_BaoCaoEmail_DanhDauDaGuiStp @NgayGui = ?", [str(ngay_gui)])
        try:
            cur.nextset()
        except Exception:
            pass
        raw.commit()
    finally:
        raw.close()


# ----------------------------------------------------------------------------- dựng PDF
def tao_pdf(path: Path, ngay, tongquan: dict, theonv: list[dict], topn: int = 15) -> None:
    from fpdf import FPDF

    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_font("DV", "", str(FONT_DIR / "DejaVuSans.ttf"))
    pdf.add_font("DV", "B", str(FONT_DIR / "DejaVuSans-Bold.ttf"))
    pdf.add_page()

    pdf.set_font("DV", "B", 15)
    pdf.cell(0, 9, config.COMPANY, new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.set_font("DV", "", 10)
    pdf.set_text_color(90, 90, 90)
    pdf.cell(0, 6, config.COMPANY_SUB, new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.ln(2)
    pdf.set_text_color(20, 20, 20)
    pdf.set_font("DV", "B", 13)
    pdf.cell(0, 8, f"BÁO CÁO BÁN HÀNG NGÀY {dinh_dang_ngay(ngay)}", new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.ln(3)

    # --- khối tổng quan
    def dong(nhan, gia_tri, dam=False):
        pdf.set_font("DV", "B" if dam else "", 11)
        pdf.set_fill_color(240, 244, 250)
        pdf.cell(95, 8, "  " + nhan, border=1, fill=True)
        pdf.set_font("DV", "B", 11)
        pdf.cell(0, 8, tien(gia_tri) + "  ", border=1, new_x="LMARGIN", new_y="NEXT", align="R")

    pdf.set_font("DV", "B", 11)
    pdf.set_fill_color(29, 95, 191)
    pdf.set_text_color(255, 255, 255)
    pdf.cell(0, 8, "  TỔNG QUAN", border=1, fill=True, new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(20, 20, 20)
    dong("Doanh số bán (đã duyệt/gửi kho)", tongquan.get("DoanhSoBan"), dam=True)
    dong("Số phiếu bán", tongquan.get("SoPhieuBan"))
    dong("Số khách mua", tongquan.get("SoKhach"))
    dong("Tiền thu tại điểm bán", tongquan.get("ThuTaiDiemBan"))
    dong("Thu công nợ (phiếu thu)", tongquan.get("ThuCongNo"))
    dong("Tổng tiền thu trong ngày", tongquan.get("TongThu"), dam=True)
    dong("Công nợ phát sinh trong ngày", tongquan.get("NoPhatSinh"), dam=True)
    # đơn web hiển thị riêng (không phải tiền)
    pdf.set_font("DV", "", 11)
    pdf.set_fill_color(240, 244, 250)
    pdf.cell(95, 8, "  Đơn web sỉ / lẻ trong ngày", border=1, fill=True)
    pdf.set_font("DV", "B", 11)
    pdf.cell(0, 8, f"{tongquan.get('DonSi', 0)} / {tongquan.get('DonLe', 0)}  ", border=1,
             new_x="LMARGIN", new_y="NEXT", align="R")
    pdf.ln(4)

    # --- bảng theo nhân viên
    pdf.set_font("DV", "B", 11)
    pdf.set_fill_color(29, 95, 191)
    pdf.set_text_color(255, 255, 255)
    pdf.cell(0, 8, f"  DOANH SỐ & CÔNG NỢ THEO NHÂN VIÊN (Top {min(topn, len(theonv))})",
             border=1, fill=True, new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(20, 20, 20)

    w = [10, 62, 20, 34, 32, 32]  # STT, Tên NV, Số phiếu, Doanh số, Tiền thu, Nợ PS
    heads = ["TT", "Nhân viên", "Phiếu", "Doanh số", "Tiền thu", "Nợ p.sinh"]
    pdf.set_font("DV", "B", 9.5)
    pdf.set_fill_color(225, 232, 242)
    for i, h in enumerate(heads):
        pdf.cell(w[i], 7, h, border=1, fill=True, align="C")
    pdf.ln()

    pdf.set_font("DV", "", 9.5)
    tong_ds = tong_thu = tong_no = 0
    for idx, nv in enumerate(theonv[:topn], 1):
        ten = str(nv.get("EmployeeName") or nv.get("EmployeeID") or "")
        if len(ten) > 34:
            ten = ten[:33] + "…"
        pdf.cell(w[0], 7, str(idx), border=1, align="C")
        pdf.cell(w[1], 7, " " + ten, border=1)
        pdf.cell(w[2], 7, str(nv.get("SoPhieu", 0)), border=1, align="C")
        pdf.cell(w[3], 7, tien(nv.get("DoanhSo")) + " ", border=1, align="R")
        pdf.cell(w[4], 7, tien(nv.get("TienThu")) + " ", border=1, align="R")
        pdf.cell(w[5], 7, tien(nv.get("NoPhatSinh")) + " ", border=1, align="R", new_x="LMARGIN", new_y="NEXT")
    # tổng cộng tất cả NV (không chỉ top)
    for nv in theonv:
        tong_ds += float(nv.get("DoanhSo") or 0)
        tong_thu += float(nv.get("TienThu") or 0)
        tong_no += float(nv.get("NoPhatSinh") or 0)
    pdf.set_font("DV", "B", 9.5)
    pdf.set_fill_color(240, 244, 250)
    pdf.cell(w[0] + w[1] + w[2], 7, "  TỔNG CỘNG (tất cả NV)", border=1, fill=True)
    pdf.cell(w[3], 7, tien(tong_ds) + " ", border=1, fill=True, align="R")
    pdf.cell(w[4], 7, tien(tong_thu) + " ", border=1, fill=True, align="R")
    pdf.cell(w[5], 7, tien(tong_no) + " ", border=1, fill=True, align="R", new_x="LMARGIN", new_y="NEXT")

    pdf.ln(6)
    pdf.set_font("DV", "", 8)
    pdf.set_text_color(120, 120, 120)
    pdf.cell(0, 5, f"Tạo tự động lúc {dt.datetime.now():%d/%m/%Y %H:%M} — Web Nguyên Quân",
             new_x="LMARGIN", new_y="NEXT")

    path.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(path))


# ----------------------------------------------------------------------------- dựng Excel
def tao_excel(path: Path, ngay, tongquan: dict, theonv: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Tổng quan"
    blue = PatternFill("solid", fgColor="1D5FBF")
    head = Font(bold=True, color="FFFFFF")
    bold = Font(bold=True)

    ws["A1"] = config.COMPANY
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"BÁO CÁO BÁN HÀNG NGÀY {dinh_dang_ngay(ngay)}"
    ws["A2"].font = Font(bold=True, size=12)
    r = 4
    tq = [
        ("Doanh số bán (đã duyệt/gửi kho)", tongquan.get("DoanhSoBan")),
        ("Số phiếu bán", tongquan.get("SoPhieuBan")),
        ("Số khách mua", tongquan.get("SoKhach")),
        ("Tiền thu tại điểm bán", tongquan.get("ThuTaiDiemBan")),
        ("Thu công nợ (phiếu thu)", tongquan.get("ThuCongNo")),
        ("Tổng tiền thu trong ngày", tongquan.get("TongThu")),
        ("Công nợ phát sinh trong ngày", tongquan.get("NoPhatSinh")),
        ("Đơn web sỉ", tongquan.get("DonSi")),
        ("Đơn web lẻ", tongquan.get("DonLe")),
    ]
    for nhan, gt in tq:
        ws.cell(r, 1, nhan)
        c = ws.cell(r, 2, int(gt or 0))
        c.number_format = "#,##0"
        r += 1
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 20

    ws2 = wb.create_sheet("Theo nhân viên")
    heads = ["TT", "Mã NV", "Nhân viên", "Số phiếu", "Doanh số", "Tiền thu", "Nợ phát sinh"]
    for i, h in enumerate(heads, 1):
        c = ws2.cell(1, i, h)
        c.fill = blue
        c.font = head
        c.alignment = Alignment(horizontal="center")
    for idx, nv in enumerate(theonv, 1):
        ws2.cell(idx + 1, 1, idx)
        ws2.cell(idx + 1, 2, str(nv.get("EmployeeID") or ""))
        ws2.cell(idx + 1, 3, str(nv.get("EmployeeName") or ""))
        ws2.cell(idx + 1, 4, int(nv.get("SoPhieu") or 0))
        for col, key in ((5, "DoanhSo"), (6, "TienThu"), (7, "NoPhatSinh")):
            c = ws2.cell(idx + 1, col, int(nv.get(key) or 0))
            c.number_format = "#,##0"
    n = len(theonv) + 2
    ws2.cell(n, 3, "TỔNG CỘNG").font = bold
    for col, key in ((5, "DoanhSo"), (6, "TienThu"), (7, "NoPhatSinh")):
        c = ws2.cell(n, col, int(sum(float(x.get(key) or 0) for x in theonv)))
        c.number_format = "#,##0"
        c.font = bold
    for col, wd in zip("ABCDEFG", (5, 12, 30, 9, 16, 16, 16)):
        ws2.column_dimensions[col].width = wd

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))


# ----------------------------------------------------------------------------- gửi mail
def gui_mail(to_list, cc_list, subject, body, attachments) -> None:
    host = config.get("email", "host", "")
    port = int(config.get("email", "port", "587") or "587")
    user = config.get("email", "user", "")
    pwd = config.get("email", "password", "")
    from_addr = config.get("email", "from", user)
    use_ssl = config.get_bool("email", "use_ssl", False)
    use_tls = config.get_bool("email", "use_tls", True)
    if not host or not from_addr:
        raise RuntimeError("Chưa cấu hình SMTP trong [email] của config.ini (host, from, user, password).")

    msg = EmailMessage()
    msg["From"] = from_addr
    msg["To"] = ", ".join(to_list)
    if cc_list:
        msg["Cc"] = ", ".join(cc_list)
    msg["Subject"] = subject
    msg.set_content(body)
    for fn, data, mime in attachments:
        maintype, subtype = mime.split("/", 1)
        msg.add_attachment(data, maintype=maintype, subtype=subtype, filename=fn)

    all_rcpt = to_list + cc_list
    if use_ssl:
        ctx = ssl.create_default_context()
        with smtplib.SMTP_SSL(host, port, context=ctx, timeout=60) as s:
            if user:
                s.login(user, pwd)
            s.send_message(msg, from_addr=from_addr, to_addrs=all_rcpt)
    else:
        with smtplib.SMTP(host, port, timeout=60) as s:
            s.ehlo()
            if use_tls:
                s.starttls(context=ssl.create_default_context())
                s.ehlo()
            if user:
                s.login(user, pwd)
            s.send_message(msg, from_addr=from_addr, to_addrs=all_rcpt)


# ----------------------------------------------------------------------------- main
def main() -> int:
    ap = argparse.ArgumentParser(description="Gửi email báo cáo bán hàng cuối ngày")
    ap.add_argument("--force", action="store_true", help="Gửi ngay, bỏ qua canh giờ")
    ap.add_argument("--ngay", help="Báo cáo cho ngày cụ thể YYYY-MM-DD (không đánh dấu đã gửi)")
    ap.add_argument("--preview", action="store_true", help="Chỉ tạo file, không gửi mail")
    ap.add_argument("--test-email", action="store_true", help="Gửi 1 mail thử để kiểm tra SMTP")
    args = ap.parse_args()

    erp = Erp()

    if args.test_email:
        cfg = kiem_tra(erp)
        to = tach_email(cfg.get("EmailNhan", ""))
        if not to:
            print("Chưa có EmailNhan trong bảng NQW_BaoCaoEmailCf.")
            return 1
        gui_mail(to, [], "[Thử] Web Nguyên Quân - kiểm tra gửi mail",
                 "Đây là email thử nghiệm cấu hình SMTP. Nếu Sếp nhận được thì cấu hình đã đúng.", [])
        print("Đã gửi mail thử tới:", ", ".join(to))
        return 0

    # xác định ngày báo cáo + có gửi hay không
    manual = bool(args.force or args.ngay or args.preview)
    cfg = kiem_tra(erp)
    if not manual:
        if not cfg.get("CanGui"):
            print(f"[{dt.datetime.now():%Y-%m-%d %H:%M}] Chưa gửi — {cfg.get('LyDo')}")
            return 0

    if args.ngay:
        ngay = dt.date.fromisoformat(args.ngay)
    else:
        ngay = cfg.get("Ngay") or dt.date.today()
        if isinstance(ngay, dt.datetime):
            ngay = ngay.date()

    tongquan, theonv = lay_du_lieu(erp, ngay)
    if not tongquan:
        print("Không lấy được dữ liệu báo cáo.")
        return 1

    dinh_dang = (cfg.get("DinhDang") or "PDF").upper()
    ngay_str = ngay.strftime("%Y%m%d") if isinstance(ngay, dt.date) else str(ngay)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    attachments = []
    files = []
    if "PDF" in dinh_dang:
        p = OUT_DIR / f"BaoCaoBanHang_{ngay_str}.pdf"
        tao_pdf(p, ngay, tongquan, theonv)
        files.append(p)
        attachments.append((p.name, p.read_bytes(), "application/pdf"))
    if "EXCEL" in dinh_dang:
        p = OUT_DIR / f"BaoCaoBanHang_{ngay_str}.xlsx"
        tao_excel(p, ngay, tongquan, theonv)
        files.append(p)
        attachments.append((p.name, p.read_bytes(),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))

    if args.preview:
        print("Đã tạo file (không gửi):")
        for f in files:
            print("  ", f)
        return 0

    tieude = (cfg.get("TieuDe") or "Báo cáo bán hàng ngày {Ngay}").replace("{Ngay}", dinh_dang_ngay(ngay))
    body = (
        f"Kính gửi Sếp,\n\n"
        f"Báo cáo bán hàng ngày {dinh_dang_ngay(ngay)} của {config.COMPANY}:\n"
        f"  - Doanh số bán: {tien(tongquan.get('DoanhSoBan'))} đ ({tongquan.get('SoPhieuBan')} phiếu, "
        f"{tongquan.get('SoKhach')} khách)\n"
        f"  - Tổng tiền thu: {tien(tongquan.get('TongThu'))} đ\n"
        f"  - Công nợ phát sinh: {tien(tongquan.get('NoPhatSinh'))} đ\n\n"
        f"Chi tiết trong file đính kèm.\n\n"
        f"(Email tự động từ Web Nguyên Quân)"
    )
    to = tach_email(cfg.get("EmailNhan", ""))
    cc = tach_email(cfg.get("EmailCC", "") or "")
    if not to:
        print("Chưa có EmailNhan trong bảng NQW_BaoCaoEmailCf.")
        return 1
    gui_mail(to, cc, tieude, body, attachments)
    print(f"[{dt.datetime.now():%Y-%m-%d %H:%M}] Đã gửi báo cáo ngày {dinh_dang_ngay(ngay)} tới {', '.join(to)}")

    # đánh dấu đã gửi hôm nay (chỉ khi là luồng gửi bình thường/force cho hôm nay)
    if not args.ngay:
        danh_dau_da_gui(erp, dt.date.today())
    return 0


if __name__ == "__main__":
    sys.exit(main())
