"""Web quản lý Nguyên Quân – bản viết lại, đọc dữ liệu NppNguyenQuan theo cấu hình web cũ."""
from __future__ import annotations

import io
import re
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from urllib.parse import quote

from fastapi import Depends, FastAPI, HTTPException, Request
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from starlette.middleware.sessions import SessionMiddleware

from . import appdb, config
from .erp import FALLBACK_CAPTIONS, Erp, ErpError, _to_date

STATIC = Path(__file__).parent / "static"
app = FastAPI(title="Nguyên Quân Web", docs_url=None, redoc_url=None)
app.add_middleware(SessionMiddleware, secret_key=config.secret_key(), max_age=12 * 3600, same_site="lax")
_erp: Erp | None = None
BUILTIN_FORMS = {"NguoiDungFrm", "NguoiDungPhanQuyenFrm"}   # màn hình hệ thống do web này tự làm


def erp() -> Erp:
    global _erp
    if _erp is None:
        _erp = Erp()
    return _erp


@app.on_event("startup")
def _startup():
    appdb.init()


@app.exception_handler(ErpError)
def _erp_error(_: Request, exc: ErpError):
    return JSONResponse({"detail": str(exc)}, status_code=400)


@app.exception_handler(PermissionError)
def _perm_error(_: Request, exc: PermissionError):
    return JSONResponse({"detail": str(exc)}, status_code=403)


@app.exception_handler(ValueError)
def _value_error(_: Request, exc: ValueError):
    return JSONResponse({"detail": str(exc)}, status_code=400)


@app.exception_handler(Exception)
def _any_error(_: Request, exc: Exception):
    msg = str(exc).splitlines()[0][:400] if str(exc) else exc.__class__.__name__
    return JSONResponse({"detail": "Lỗi đọc dữ liệu: " + msg}, status_code=500)


# ---------------------------------------------------------------------- đăng nhập
def _admin_groups() -> set[str]:
    g = config.get("app", "bms_admin_groups", "Admin,Admin1,Admin2,Admin3")
    return {x.strip().lower() for x in g.split(",") if x.strip()}


def _from_sy_user(row: dict) -> dict:
    """Chuyển 1 dòng SY_User thành 'user' của web (đăng nhập bằng tài khoản phần mềm)."""
    nhom = row.get("UserGroupID")
    admin = bool(nhom) and nhom.strip().lower() in _admin_groups()
    return {"ten_dang_nhap": row["UserName"], "ho_ten": row.get("HoTen") or row.get("TenNgan") or row["UserName"],
            "la_admin": admin, "nhom": nhom, "ma_nv": (row.get("EmployeeID") or None) if not admin else None,
            "dang_dung": not row.get("Disable"), "mat_khau": None, "_src": "sy_user"}


def current_user(request: Request) -> dict:
    name = request.session.get("u")
    if not name:
        raise HTTPException(401, "Chưa đăng nhập")
    if request.session.get("src") == "sy_user":
        row = erp().sy_user(name)
        if not row or row.get("Disable"):
            raise HTTPException(401, "Tài khoản đã bị khóa hoặc không còn")
        return _from_sy_user(row)
    u = appdb.get_user(name)
    if not u or not u["dang_dung"]:
        raise HTTPException(401, "Chưa đăng nhập")
    return u


def admin_user(u: dict = Depends(current_user)) -> dict:
    if not u["la_admin"]:
        raise HTTPException(403, "Chỉ quản trị mới dùng được chức năng này")
    return u


class LoginIn(BaseModel):
    username: str
    password: str


class PasswordIn(BaseModel):
    old: str
    new: str


@app.post("/api/login")
def login(body: LoginIn, request: Request):
    name = body.username.strip()
    # 0) chặn nếu tài khoản đang bị khóa do nhập sai nhiều lần
    khoa = appdb.kiem_khoa(name)
    if khoa["bi_khoa"]:
        appdb.log(name[:50], "Đăng nhập bị chặn", "đang khóa")
        raise HTTPException(423, khoa["thong_bao"])
    # 1) tài khoản riêng của web (quản trị web)
    u = appdb.get_user(name)
    if u and u["dang_dung"] and appdb.check_password(body.password, u["mat_khau"]):
        appdb.reset_that_bai(name)
        request.session.update({"u": u["ten_dang_nhap"], "src": "local"})
        appdb.log(u["ten_dang_nhap"], "Đăng nhập")
        return {"ok": True}
    # 2) tài khoản nhân viên trong phần mềm (SY_User) — băm lại theo thuật toán phần mềm
    try:
        row = erp().verify_sy_user(name, body.password)
    except Exception:
        row = None
    if row:
        appdb.reset_that_bai(name)
        request.session.update({"u": row["UserName"], "src": "sy_user"})
        appdb.log(row["UserName"], "Đăng nhập (SY_User)")
        return {"ok": True}
    # 3) sai mật khẩu -> ghi nhận, áp khóa nếu tới ngưỡng
    st = appdb.ghi_that_bai(name)
    appdb.log(name[:50], "Đăng nhập sai", f"lần thứ {st['so_lan_sai']}")
    khoa = appdb.kiem_khoa(name)
    if khoa["bi_khoa"]:
        raise HTTPException(423, khoa["thong_bao"])
    con_lai = appdb.KHOA_NGUONG_1 - st["so_lan_sai"]
    msg = "Sai tên đăng nhập hoặc mật khẩu."
    if 0 < con_lai <= 2:
        msg += f" Còn {con_lai} lần trước khi tài khoản bị tạm khóa."
    raise HTTPException(401, msg)


@app.post("/api/logout")
def logout(request: Request):
    request.session.clear()
    return {"ok": True}


# -------- Khóa đăng nhập: admin xem danh sách & mở khóa --------
@app.get("/api/khoa/danh-sach")
def khoa_danh_sach(u: dict = Depends(admin_user)):
    return appdb.ds_khoa()


@app.post("/api/khoa/mo/{name}")
def khoa_mo(name: str, u: dict = Depends(admin_user)):
    appdb.mo_khoa(name)
    appdb.log(u["ten_dang_nhap"], "Mở khóa đăng nhập", name)
    return {"ok": True}


@app.post("/api/doi-mat-khau")
def change_password(body: PasswordIn, u: dict = Depends(current_user)):
    if u.get("_src") == "sy_user":
        raise HTTPException(400, "Tài khoản nhân viên đổi mật khẩu trong phần mềm bán hàng, không đổi ở web")
    if not appdb.check_password(body.old, u["mat_khau"]):
        raise HTTPException(400, "Mật khẩu cũ không đúng")
    if len(body.new) < 6:
        raise HTTPException(400, "Mật khẩu mới tối thiểu 6 ký tự")
    appdb.set_password(u["ten_dang_nhap"], body.new)
    appdb.log(u["ten_dang_nhap"], "Đổi mật khẩu")
    return {"ok": True}


# ---------------------------------------------------------------------- menu + quyền
def _perm(u: dict, menu_id: str) -> dict:
    """Trả về menu nếu người dùng được chạy, kèm quyền xuất Excel."""
    m = erp().menu_item(menu_id)
    if m.get("FormName") in BUILTIN_FORMS:
        raise HTTPException(404, "Màn hình này nằm trong mục Người dùng web")
    if u["la_admin"]:
        if m.get("isDisable") and not config.SHOW_DISABLED_MENU:
            raise HTTPException(403, "Menu đang tắt")
        return {**m, "excel": True, "add": True, "edit": True, "delete": True}
    if m.get("isDisable"):
        raise HTTPException(403, "Menu đang tắt")
    p = erp().group_permissions(u["nhom"] or "").get(str(menu_id)) if u["nhom"] else None
    if not p or not p.get("IsRun"):
        raise HTTPException(403, "Bạn chưa được cấp quyền dùng màn hình này")
    return {**m, "excel": bool(p.get("isExportExcel")), "add": bool(p.get("IsAdd")),
            "edit": bool(p.get("IsUpdate")), "delete": bool(p.get("IsDelete"))}


def _entry(m: dict) -> dict | None:
    """Cấu hình nhập đơn (sy_WebLoaiDon) nếu màn hình này là đơn hàng sỉ/lẻ."""
    t = erp().order_types().get(m.get("FormName") or "")
    if not t:
        return None
    return {"loai": t["LoaiDon"], "ten": t["TenLoai"], "ck": bool(t["ChoPhepChietKhau"]),
            "add": m["add"], "edit": m["edit"], "delete": m["delete"]}


def _menu_tree(u: dict, skipped: list | None = None) -> list[dict]:
    items = erp().menu_all()
    perms = erp().group_permissions(u["nhom"]) if (u["nhom"] and not u["la_admin"]) else {}
    leaves = []
    for m in items:
        if not m.get("Parent") or m.get("FormName") in BUILTIN_FORMS or m.get("FormKey") not in ("List", "Report"):
            continue
        if m.get("isDisable") and not (u["la_admin"] and config.SHOW_DISABLED_MENU):
            continue
        if not u["la_admin"]:
            p = perms.get(str(m["MenuID"]))
            if not p or not p.get("IsRun"):
                continue
        why = erp().target_check(m)
        if why:
            if skipped is not None:
                skipped.append({"id": str(m["MenuID"]), "title": m["VN"], "why": why})
            continue
        leaves.append(m)
    groups = []
    for g in items:
        kids = [{"id": str(k["MenuID"]), "title": k["VN"], "kind": k["FormKey"].lower(), "off": bool(k.get("isDisable"))}
                for k in leaves if str(k["Parent"]) == str(g["MenuID"])]
        if kids:
            groups.append({"id": str(g["MenuID"]), "title": g["VN"], "items": kids})
    return groups


@app.get("/api/me")
def me(u: dict = Depends(current_user)):
    skipped: list = []
    try:
        tree, err = _menu_tree(u, skipped), None
    except Exception as e:  # vẫn cho vào web để admin sửa cấu hình
        tree, err = [], "Không đọc được menu từ database: " + str(e).splitlines()[0][:300]
    return {"user": u["ten_dang_nhap"], "name": u["ho_ten"], "admin": bool(u["la_admin"]), "ma_nv": u["ma_nv"],
            "nguon": u.get("_src", "local"), "menu": tree, "menu_error": err,
            "menu_skipped": skipped if u["la_admin"] else [], "company": config.COMPANY, "company_sub": config.COMPANY_SUB}


# ---------------------------------------------------------------------- dữ liệu -> JSON
def _jv(v):
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%dT%H:%M:%S")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (bytes, bytearray)):
        return None
    return v


def _rows_json(rows: list[dict], cols: list[str]) -> list[list]:
    return [[_jv(r.get(c)) for c in cols] for r in rows]


# ---------------------------------------------------------------------- trang chủ
@app.get("/api/dashboard")
def dashboard(den: str | None = None, u: dict = Depends(current_user)):
    d = erp().dashboard(_to_date(den), u["ma_nv"])
    return {k: ({kk: _jv(vv) for kk, vv in v.items()} if isinstance(v, dict) else
                [{kk: _jv(vv) for kk, vv in r.items()} for r in v] if isinstance(v, list) else v) for k, v in d.items()}


# ---------------------------------------------------------------------- màn hình danh sách
class ListQuery(BaseModel):
    tu: str | None = None
    den: str | None = None
    filters: dict[str, str] = {}
    q: str = ""


@app.get("/api/list/{menu_id}/meta")
def list_meta(menu_id: str, u: dict = Depends(current_user)):
    m = _perm(u, menu_id)
    d = erp().form_def(m["FormName"])
    fm = erp().formats(m["FormName"])
    filters = []
    for f in d["filters"]:
        locked = bool(u["ma_nv"]) and f["col"].lower() == "employeeid"
        default = u["ma_nv"] if locked else ""
        if not locked and f["default"] == "{EmployeeID}":
            default = ""
        filters.append({"col": f["col"], "caption": f["caption"] or (fm.get(f["col"].lower()) or {}).get("caption") or FALLBACK_CAPTIONS.get(f["col"].lower()) or f["col"],
                        "has_options": bool(f["source"]), "default": default, "locked": locked})
    today = date.today()
    return {"title": m["VN"], "excel": m["excel"], "has_date": bool(d["master"]["date"]),
            "date_caption": (fm.get((d["master"]["date"] or "").lower()) or {}).get("caption") or "Ngày",
            "tu": today.isoformat(), "den": today.isoformat(),
            "filters": filters, "has_detail": bool(d["detail"]), "entry": _entry(m)}


@app.get("/api/list/{menu_id}/options/{col}")
def list_options(menu_id: str, col: str, u: dict = Depends(current_user)):
    m = _perm(u, menu_id)
    only = u["ma_nv"] if (u["ma_nv"] and col.lower() == "employeeid") else None
    return erp().filter_options(m["FormName"], col, only)


def _list_data(menu_id: str, body: ListQuery, u: dict):
    m = _perm(u, menu_id)
    d = erp().form_def(m["FormName"])
    rows, truncated = erp().list_rows(m["FormName"], _to_date(body.tu), _to_date(body.den), body.filters, body.q, u["ma_nv"])
    mm = d["master"]
    desc = erp().describe(mm["grid"], m["FormName"], mm["sums"], rows)
    return m, d, rows, desc, truncated


@app.post("/api/list/{menu_id}/rows")
def list_rows(menu_id: str, body: ListQuery, u: dict = Depends(current_user)):
    m, d, rows, desc, truncated = _list_data(menu_id, body, u)
    mm = d["master"]
    extra = [mm["pk"]] + [c for c in (mm["status"], mm["color"]) if c]
    return {"cols": desc, "rows": _rows_json(rows, [c["name"] for c in desc]), "keys": _rows_json(rows, extra),
            "key_names": {"pk": 0, "status": 1 if mm["status"] else None, "color": (2 if mm["status"] else 1) if mm["color"] else None},
            "truncated": truncated, "max_rows": config.MAX_ROWS}


@app.get("/api/list/{menu_id}/record/{key:path}")
def list_record(menu_id: str, key: str, u: dict = Depends(current_user)):
    m = _perm(u, menu_id)
    d = erp().form_def(m["FormName"])
    head, lines, dcols = erp().record(m["FormName"], key, u["ma_nv"])
    fm = erp().formats(m["FormName"])
    layout = d["layout"] or [{"name": c, "type": "text"} for c in d["master"]["grid"]]
    layout = [f for f in layout if not (isinstance(head.get(f["name"]), str) and "<a " in head.get(f["name"]).lower())]
    fields = [{"name": f["name"], "caption": (fm.get(f["name"].lower()) or {}).get("caption") or FALLBACK_CAPTIONS.get(f["name"].lower()) or f["name"], "type": f["type"],
               "value": _jv(head.get(f["name"]))} for f in layout if f["name"] in head]
    det = d["detail"] or {}
    dcols = [c for c in dcols if c.lower() not in ("latitude", "longitude")]
    desc = erp().describe(dcols, m["FormName"], det.get("sums"), lines)
    lat = next((c for c in (lines[0].keys() if lines else []) if c.lower() == "latitude"), None)
    lng = next((c for c in (lines[0].keys() if lines else []) if c.lower() == "longitude"), None)
    points = []
    if lat and lng:
        name_col = next((c for c in lines[0].keys() if c.lower() == "objectname"), None)
        addr_col = next((c for c in lines[0].keys() if c.lower() == "address"), None)
        for i, r in enumerate(lines):
            try:
                la, ln = float(r[lat]), float(r[lng])
            except (TypeError, ValueError):
                continue
            if la and ln:
                points.append({"i": i + 1, "lat": la, "lng": ln, "name": r.get(name_col) if name_col else "",
                               "addr": r.get(addr_col) if addr_col else ""})
    appdb.log(u["ten_dang_nhap"], "Xem chứng từ", f"{m['VN']}: {key}")
    ent = _entry(m)
    st = str(head.get("StatusID", ""))
    lock = "Đã duyệt – không sửa được" if (head.get("isSelect") or st == "1") else \
        "Chờ duyệt – không sửa được (giống web cũ)" if st == "0" else None
    return {"title": m["VN"], "key": key, "entry": ent, "approved": bool(lock), "lock_reason": lock, "fields": fields, "cols": desc, "lines": _rows_json(lines, dcols), "points": points}


@app.post("/api/list/{menu_id}/excel")
def list_excel(menu_id: str, body: ListQuery, u: dict = Depends(current_user)):
    m, d, rows, desc, _ = _list_data(menu_id, body, u)
    if not m["excel"]:
        raise HTTPException(403, "Bạn chưa được cấp quyền xuất Excel")
    sub = f"Từ {body.tu or ''} đến {body.den or ''}" if d["master"]["date"] else ""
    appdb.log(u["ten_dang_nhap"], "Xuất Excel", m["VN"])
    return _xlsx(m["VN"], sub, desc, rows, None)


# ---------------------------------------------------------------------- nhập đơn hàng sỉ / lẻ
class OrderLine(BaseModel):
    ItemID: str
    Ket: float | None = None
    Chai: float | None = None
    KetKM: float | None = None
    ChaiKM: float | None = None
    UnitPrice: float | None = None
    TienChietKhau: float | None = None
    KhuyenMaiID: str | None = None
    Notes: str | None = None


class OrderIn(BaseModel):
    DocumentID: str = ""
    DocumentDate: str
    DeliveryDate: str | None = None
    ObjectID: str
    EmployeeID: str
    Phone: str | None = None
    Memo: str | None = None
    isTrongNgay: bool = False
    isKhongKM: bool = False
    HoanTat: bool = False
    lines: list[OrderLine] = []


def _order_ctx(u: dict, menu_id: str, need: str | None = None) -> tuple[dict, dict]:
    m = _perm(u, menu_id)
    ent = _entry(m)
    if not ent:
        raise HTTPException(400, "Màn hình này không nhập đơn được (chưa chạy file sql\\NQW_CapNhatDatabase.sql?)")
    if need and not ent[need]:
        raise HTTPException(403, {"add": "Bạn chưa được cấp quyền thêm đơn",
                                  "edit": "Bạn chưa được cấp quyền sửa đơn",
                                  "delete": "Bạn chưa được cấp quyền xóa đơn"}[need])
    return m, ent


def _emp_for(u: dict, emp: str | None) -> str:
    return u["ma_nv"] if u["ma_nv"] else (emp or "").strip()


def _rows_out(rows: list[dict]) -> list[dict]:
    return [{k: _jv(v) for k, v in r.items()} for r in rows]


@app.get("/api/order/{menu_id}/employees")
def order_employees(menu_id: str, u: dict = Depends(current_user)):
    _order_ctx(u, menu_id)
    emps = erp().employees()
    if u["ma_nv"]:
        emps = [e for e in emps if e["EmployeeID"] == u["ma_nv"]]
    return [{"value": e["EmployeeID"], "label": f"{e['EmployeeName']} - {e['EmployeeID']}"} for e in emps]


@app.get("/api/order/{menu_id}/customers")
def order_customers(menu_id: str, emp: str = "", q: str = "", u: dict = Depends(current_user)):
    _, ent = _order_ctx(u, menu_id)
    rows = erp().call_proc("NQW_DonHang_KhachHangStp", [("LoaiDon", ent["loai"]), ("EmployeeID", _emp_for(u, emp)), ("S", q.strip()[:100])])
    return _rows_out(rows)


@app.get("/api/order/{menu_id}/customer/{obj}")
def order_customer(menu_id: str, obj: str, ngay: str | None = None, u: dict = Depends(current_user)):
    _, ent = _order_ctx(u, menu_id)
    rows = erp().call_proc("NQW_DonHang_ThongTinKHStp", [("ObjectID", obj), ("Ngay", _to_date(ngay) or date.today())])
    if not rows:
        raise HTTPException(404, "Không tìm thấy khách hàng " + obj)
    r = rows[0]
    if u["ma_nv"] and r.get("EmployeeID") != u["ma_nv"]:
        raise HTTPException(403, "Khách hàng không do bạn quản lý")
    return {k: _jv(v) for k, v in r.items()}


@app.get("/api/order/{menu_id}/items")
def order_items(menu_id: str, obj: str = "", u: dict = Depends(current_user)):
    _, ent = _order_ctx(u, menu_id)
    return _rows_out(erp().call_proc("NQW_DonHang_HangHoaStp", [("LoaiDon", ent["loai"]), ("ObjectID", obj)]))


@app.get("/api/order/{menu_id}/promos")
def order_promos(menu_id: str, ngay: str | None = None, u: dict = Depends(current_user)):
    """Chương trình khuyến mãi còn hiệu lực (giống ô CTKM của web cũ: B_KhuyenMaiTheoThoiGianFnc)."""
    _order_ctx(u, menu_id)
    e = erp()
    if not e.object_exists("B_KhuyenMaiTheoThoiGianFnc"):
        return []
    d = _to_date(ngay) or date.today()
    rows = e.rows("SELECT KhuyenMaiID, TenCTKM FROM dbo.B_KhuyenMaiTheoThoiGianFnc(:d, N'')", {"d": datetime.combine(d, datetime.min.time())})
    return [{"value": str(r["KhuyenMaiID"]), "label": f"{r['TenCTKM']} - {r['KhuyenMaiID']}"} for r in rows if r.get("KhuyenMaiID")]


@app.get("/api/order/{menu_id}/doc/{key:path}")
def order_doc(menu_id: str, key: str, u: dict = Depends(current_user)):
    _, ent = _order_ctx(u, menu_id)
    head, lines = erp().order_doc(key)
    if not head:
        raise HTTPException(404, "Không tìm thấy đơn " + key)
    if u["ma_nv"] and head.get("EmployeeID") != u["ma_nv"]:
        raise HTTPException(403, "Bạn chỉ được xem đơn của nhân viên " + u["ma_nv"])
    keep = ["ItemID", "ItemName", "Unit", "QuyDoi", "Ket", "Chai", "KetKM", "ChaiKM", "GiaChuan", "UnitPrice", "TienChietKhau",
            "KhuyenMaiID", "Notes", "Amount"]
    return {"head": {k: _jv(v) for k, v in head.items()}, "lines": [{k: _jv(l.get(k)) for k in keep} for l in lines],
            "approved": bool(head.get("isSelect")) or (head.get("MaTrangThai") not in (None, -1))}


def _xml_lines(lines: list[OrderLine]) -> str:
    from xml.sax.saxutils import escape
    parts = []
    for ln in lines:
        cells = []
        for k, v in ln.model_dump().items():
            if v is None or v == "":
                continue
            if isinstance(v, float):
                v = int(v) if float(v).is_integer() else round(v)
            cells.append(f"<{k}>{escape(str(v).strip())}</{k}>")
        parts.append("<L>" + "".join(cells) + "</L>")
    return "<Lines>" + "".join(parts) + "</Lines>"


def _messages(rows: list[dict]) -> tuple[list[str], list[str], str | None]:
    errors = [str(r["Title"]) for r in rows if r.get("MsgType") == 1]
    warns = [str(r["Title"]) for r in rows if r.get("MsgType") == 0]
    doc = next((r.get("DocumentID") for r in rows if r.get("DocumentID")), None)
    return errors, warns, doc


@app.post("/api/order/{menu_id}/save")
def order_save(menu_id: str, body: OrderIn, u: dict = Depends(current_user)):
    new = not body.DocumentID.strip()
    m, ent = _order_ctx(u, menu_id, "add" if new else "edit")
    emp = _emp_for(u, body.EmployeeID)
    if not emp:
        raise HTTPException(400, "Chọn nhân viên bán hàng")
    if not body.ObjectID.strip():
        raise HTTPException(400, "Chọn khách hàng")
    lines = [ln for ln in body.lines if ln.ItemID.strip()]
    if not lines:
        raise HTTPException(400, "Đơn hàng chưa có mặt hàng nào")
    if not new:
        head, _ = erp().order_doc(body.DocumentID.strip())
        if not head:
            raise HTTPException(404, "Không tìm thấy đơn " + body.DocumentID)
        if u["ma_nv"] and head.get("EmployeeID") != u["ma_nv"]:
            raise HTTPException(403, "Bạn chỉ được sửa đơn của nhân viên " + u["ma_nv"])
    args = [("DocumentID", body.DocumentID.strip()), ("LoaiDon", ent["loai"]), ("DocumentDate", _to_date(body.DocumentDate)),
            ("DeliveryDate", _to_date(body.DeliveryDate) or _to_date(body.DocumentDate)), ("ObjectID", body.ObjectID.strip()),
            ("EmployeeID", emp), ("Phone", (body.Phone or "").strip()[:50] or None), ("Memo", (body.Memo or "").strip()[:200] or None),
            ("isTrongNgay", 1 if body.isTrongNgay else 0), ("isKhongKM", 1 if body.isKhongKM else 0), ("HoanTat", 1 if body.HoanTat else 0),
            ("Lines", _xml_lines(lines)), ("User", u["ten_dang_nhap"][:20]), ("IsAdmin", 1 if u["la_admin"] else 0)]
    rows = erp().call_proc("NQW_DonHang_LuuStp", args, write=True)
    errors, warns, doc = _messages(rows)
    if errors or not doc:
        return JSONResponse({"ok": False, "errors": errors or ["Không lưu được đơn hàng"], "warnings": warns}, status_code=422)
    appdb.log(u["ten_dang_nhap"], "Thêm đơn" if new else "Sửa đơn", f"{m['VN']}: {doc}")
    return {"ok": True, "DocumentID": doc, "warnings": warns}


@app.delete("/api/order/{menu_id}/doc/{key:path}")
def order_delete(menu_id: str, key: str, u: dict = Depends(current_user)):
    m, _ = _order_ctx(u, menu_id, "delete")
    head, _ = erp().order_doc(key)
    if not head:
        raise HTTPException(404, "Không tìm thấy đơn " + key)
    if u["ma_nv"] and head.get("EmployeeID") != u["ma_nv"]:
        raise HTTPException(403, "Bạn chỉ được xóa đơn của nhân viên " + u["ma_nv"])
    rows = erp().call_proc("NQW_DonHang_XoaStp", [("DocumentID", key), ("User", u["ten_dang_nhap"][:20])], write=True)
    errors, _, _ = _messages(rows)
    if errors:
        raise HTTPException(400, "; ".join(errors))
    appdb.log(u["ten_dang_nhap"], "Xóa đơn", f"{m['VN']}: {key}")
    return {"ok": True}


# ---------------------------------------------------------------------- báo cáo
class ReportRun(BaseModel):
    values: dict[str, str | None] = {}


@app.get("/api/report/{menu_id}/meta")
def report_meta(menu_id: str, u: dict = Depends(current_user)):
    m = _perm(u, menu_id)
    ps = erp().proc_params(m["FormName"])
    today = date.today()
    out = []
    for p in ps:
        n = p["name"].lower()
        locked = bool(u["ma_nv"]) and n in ("employeeid", "semployeeid")
        default = u["ma_nv"] if locked else ""
        if p["type"] == "date":
            default = today.replace(day=1).isoformat() if n.startswith(("from", "tu")) else today.isoformat()
        out.append({**p, "default": default, "locked": locked})
    return {"title": m["VN"], "excel": m["excel"], "params": out}


@app.get("/api/report/{menu_id}/options/{param}")
def report_options(menu_id: str, param: str, u: dict = Depends(current_user)):
    m = _perm(u, menu_id)
    only = u["ma_nv"] if (u["ma_nv"] and param.lower() in ("employeeid", "semployeeid")) else None
    return erp().param_options(m["FormName"], param, only)


def _group_col(cols: list[str], rows: list[dict]) -> str | None:
    """Báo cáo web cũ nhóm theo cột đầu nếu giá trị dạng 'NVBH : Tên ...'."""
    if not cols or not rows:
        return None
    first = cols[0]
    vals = [r.get(first) for r in rows[:500]]
    if all(isinstance(v, str) and re.match(r"^[^:]{1,40} : ", v) for v in vals if v is not None) and any(vals):
        return first
    return None


def _report_data(menu_id: str, body: ReportRun, u: dict):
    m = _perm(u, menu_id)
    cols, rows, truncated = erp().run_report(m["FormName"], body.values, u["ma_nv"])
    group = _group_col(cols, rows)
    shown = [c for c in cols if c != group]
    desc = erp().describe(shown, m["FormName"], None, rows)
    return m, rows, desc, group, truncated


@app.post("/api/report/{menu_id}/run")
def report_run(menu_id: str, body: ReportRun, u: dict = Depends(current_user)):
    m, rows, desc, group, truncated = _report_data(menu_id, body, u)
    appdb.log(u["ten_dang_nhap"], "Xem báo cáo", m["VN"])
    return {"cols": desc, "rows": _rows_json(rows, [c["name"] for c in desc]),
            "groups": [_jv(r.get(group)) for r in rows] if group else None, "truncated": truncated, "max_rows": config.MAX_ROWS}


@app.post("/api/report/{menu_id}/excel")
def report_excel(menu_id: str, body: ReportRun, u: dict = Depends(current_user)):
    m, rows, desc, group, _ = _report_data(menu_id, body, u)
    if not m["excel"]:
        raise HTTPException(403, "Bạn chưa được cấp quyền xuất Excel")
    sub = ", ".join(f"{k}: {v}" for k, v in body.values.items() if v)
    appdb.log(u["ten_dang_nhap"], "Xuất Excel", m["VN"])
    return _xlsx(m["VN"], sub, desc, rows, group)


def _xlsx(title: str, sub: str, desc: list[dict], rows: list[dict], group: str | None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Du lieu"
    n = len(desc)
    ws.cell(1, 1, title).font = Font(bold=True, size=14)
    ws.cell(2, 1, sub).font = Font(italic=True, color="555555")
    head_fill = PatternFill("solid", fgColor="1D5FBF")
    thin = Side(style="thin", color="C9D1DB")
    for j, c in enumerate(desc, 1):
        cell = ws.cell(4, j, c["caption"])
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    fmts = {"N": "#,##0", "N2": "#,##0.00", "D": "dd/mm/yyyy"}
    r_i = 5
    last_group = object()
    for r in rows:
        if group and r.get(group) != last_group:
            last_group = r.get(group)
            ws.cell(r_i, 1, str(last_group or "")).font = Font(bold=True, color="1D5FBF")
            r_i += 1
        for j, c in enumerate(desc, 1):
            v = r.get(c["name"])
            if isinstance(v, Decimal):
                v = float(v)
            if isinstance(v, str):
                v = re.sub(r"<\s*/?\s*br\s*/?>", "\n", v, flags=re.I)
                v = re.sub(r"<[^>]+>", "", v)
            cell = ws.cell(r_i, j, v)
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if c["fmt"] in fmts:
                cell.number_format = fmts[c["fmt"]]
        r_i += 1
    if rows and any(c["sum"] for c in desc):
        for j, c in enumerate(desc, 1):
            if c["sum"]:
                col = get_column_letter(j)
                cell = ws.cell(r_i, j, f"=SUBTOTAL(9,{col}5:{col}{r_i - 1})")
                cell.number_format = fmts.get(c["fmt"], "#,##0")
                cell.font = Font(bold=True)
        ws.cell(r_i, 1).value = ws.cell(r_i, 1).value or "Tổng cộng"
        ws.cell(r_i, 1).font = Font(bold=True)
    for j, c in enumerate(desc, 1):
        width = max([len(str(c["caption"]))] + [len(str(r.get(c["name"]) or ""))
                                                for r in rows[:300]])
        ws.column_dimensions[get_column_letter(j)].width = min(max(10, width + 2), 50)
    ws.freeze_panes = "A5"
    if n:
        ws.auto_filter.ref = f"A4:{get_column_letter(n)}4"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = re.sub(r"[^\w\-]+", "_", title, flags=re.UNICODE).strip("_") + f"_{date.today():%Y%m%d}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"})


# ---------------------------------------------------------------------- quản trị
class UserIn(BaseModel):
    ten_dang_nhap: str
    ho_ten: str = ""
    la_admin: bool = False
    nhom: str | None = None
    ma_nv: str | None = None
    dang_dung: bool = True
    mat_khau: str | None = None


@app.get("/api/nguoi-dung")
def users(u: dict = Depends(admin_user)):
    return appdb.list_users()


@app.get("/api/danh-muc-quyen")
def perm_lists(u: dict = Depends(admin_user)):
    try:
        return {"nhom": erp().user_groups(), "nhan_vien": erp().employees()}
    except Exception as e:
        return {"nhom": [], "nhan_vien": [], "loi": str(e).splitlines()[0][:300]}


@app.post("/api/nguoi-dung")
def save_user(body: UserIn, u: dict = Depends(admin_user)):
    name = body.ten_dang_nhap.strip()
    if not re.fullmatch(r"[A-Za-z0-9_.\-]{2,40}", name):
        raise HTTPException(400, "Tên đăng nhập chỉ gồm chữ không dấu, số, dấu . _ - (2–40 ký tự)")
    if body.mat_khau and len(body.mat_khau) < 6:
        raise HTTPException(400, "Mật khẩu tối thiểu 6 ký tự")
    old = appdb.get_user(name)
    if old and old["la_admin"] and (not body.la_admin or not body.dang_dung) and appdb.count_admins(exclude=name) == 0:
        raise HTTPException(400, "Phải còn ít nhất 1 tài khoản quản trị đang dùng")
    if not body.la_admin and not body.nhom:
        raise HTTPException(400, "Chọn nhóm quyền cho người dùng (hoặc đánh dấu Quản trị)")
    appdb.save_user(body.model_dump(), body.mat_khau)
    appdb.log(u["ten_dang_nhap"], "Lưu người dùng", name)
    return {"ok": True}


@app.delete("/api/nguoi-dung/{name}")
def delete_user(name: str, u: dict = Depends(admin_user)):
    if name.lower() == u["ten_dang_nhap"].lower():
        raise HTTPException(400, "Không tự xóa tài khoản đang đăng nhập")
    old = appdb.get_user(name)
    if old and old["la_admin"] and appdb.count_admins(exclude=name) == 0:
        raise HTTPException(400, "Phải còn ít nhất 1 tài khoản quản trị")
    appdb.delete_user(name)
    appdb.log(u["ten_dang_nhap"], "Xóa người dùng", name)
    return {"ok": True}


@app.get("/api/nhat-ky")
def web_logs(tu: str, den: str, nguoi: str = "", u: dict = Depends(admin_user)):
    return appdb.logs(tu, den, nguoi)


@app.post("/api/lam-moi")
def refresh(u: dict = Depends(admin_user)):
    erp().clear_cache()
    return {"ok": True}


# ---------------------------------------------------------------------- giao diện
app.mount("/static", StaticFiles(directory=STATIC), name="static")


@app.get("/")
def index():
    return FileResponse(STATIC / "index.html", headers={"Cache-Control": "no-cache"})
